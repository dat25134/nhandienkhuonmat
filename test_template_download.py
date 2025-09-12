#!/usr/bin/env python3
"""
Script test download template Excel
"""
import sys
import requests
import os

def test_template_download():
    """Test download template Excel"""
    base_url = "http://localhost:5000"
    
    print("=== Testing Template Download ===\n")
    
    try:
        print("1. Testing API endpoint...")
        response = requests.get(f"{base_url}/api/excel/template", timeout=10)
        
        print(f"   Status: {response.status_code}")
        print(f"   Content-Type: {response.headers.get('Content-Type')}")
        print(f"   Content-Disposition: {response.headers.get('Content-Disposition')}")
        print(f"   Content-Length: {response.headers.get('Content-Length')}")
        
        if response.status_code == 200:
            print("   ✅ API response successful")
            
            # Lưu file để test
            test_file = "downloaded_template.xlsx"
            with open(test_file, 'wb') as f:
                f.write(response.content)
            
            print(f"   ✅ Template saved as: {test_file}")
            print(f"   ✅ File size: {os.path.getsize(test_file)} bytes")
            
            # Kiểm tra file có thể đọc được không
            try:
                from openpyxl import load_workbook
                wb = load_workbook(test_file)
                ws = wb.active
                print(f"   ✅ Excel file readable, rows: {ws.max_row}, cols: {ws.max_column}")
                
                # Hiển thị headers
                headers = []
                for col in range(1, ws.max_column + 1):
                    headers.append(ws.cell(row=1, column=col).value)
                print(f"   ✅ Headers: {headers}")
                
            except Exception as e:
                print(f"   ❌ Excel file not readable: {e}")
            
            # Cleanup
            os.remove(test_file)
            print("   ✅ Test file cleaned up")
            
        else:
            print(f"   ❌ API error: {response.text}")
            
    except requests.exceptions.ConnectionError:
        print("   ❌ Cannot connect to server. Make sure app is running.")
        print("   Run: python3 run_app.py")
        
    except Exception as e:
        print(f"   ❌ Error: {e}")
    
    print("\n=== Test Complete ===")

if __name__ == "__main__":
    test_template_download()
