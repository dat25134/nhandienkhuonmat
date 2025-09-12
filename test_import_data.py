#!/usr/bin/env python3
"""
Script tạo file Excel test để demo chức năng import
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_test_excel():
    """Tạo file Excel test cho import"""
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách khách"
    
    # Headers
    headers = [
        "STT",
        "Họ và tên", 
        "Số điện thoại",
        "Giới tính",
        "Công ty",
        "Bộ phận", 
        "Vị trí",
        "Ghi chú"
    ]
    
    # Thêm headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    # Dữ liệu test
    test_data = [
        [1, "Nguyễn Văn Test1", "0123456789", "Nam", "Công ty Test", "IT", "Nhân viên", "Test import"],
        [2, "Trần Thị Test2", "0987654321", "Nữ", "Công ty Demo", "HR", "Trưởng phòng", ""],
        [3, "Lê Văn Test3", "0369258147", "Nam", "Công ty Sample", "Marketing", "Chuyên viên", "Khách mới"],
        [4, "Phạm Thị Test4", "0555123456", "Nữ", "Công ty Example", "Sales", "Nhân viên", ""],
        [5, "Hoàng Văn Test5", "0777888999", "Nam", "Công ty Demo", "IT", "Kỹ sư", "VIP"],
    ]
    
    # Thêm dữ liệu
    for row_idx, row_data in enumerate(test_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            if col_idx == 1:  # STT column
                cell.alignment = Alignment(horizontal="center")
    
    # Điều chỉnh độ rộng cột
    column_widths = [8, 25, 15, 12, 20, 15, 15, 20]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    
    # Lưu file
    test_file = 'test_import_data.xlsx'
    wb.save(test_file)
    print(f"✅ Đã tạo file test: {test_file}")
    print("📋 Dữ liệu test:")
    for i, row in enumerate(test_data, 1):
        print(f"   {i}. {row[1]} - {row[2]} - {row[4]}")
    
    return test_file

if __name__ == "__main__":
    create_test_excel()
